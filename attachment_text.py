import csv
from email import policy
from email.parser import BytesParser
from html.parser import HTMLParser
import io
from pathlib import Path
import shutil
import subprocess
import tempfile
from xml.etree import ElementTree as ET
from zipfile import ZipFile

MEDIA_EXTENSIONS = {'.jpg', '.jpeg', '.png', '.gif', '.bmp', '.tif', '.tiff', '.webp',
                    '.heic', '.svg', '.mp4', '.mov', '.avi', '.mkv', '.wmv', '.mpeg',
                    '.mpg', '.webm', '.m4v'}
TEXT_EXTENSIONS = {'.txt', '.csv', '.tsv', '.json', '.xml', '.md', '.log'}
ATTACHMENT_EXTENSIONS = TEXT_EXTENSIONS | {'.pdf', '.docx', '.eml', '.msg', '.rtf',
    '.html', '.htm', '.xlsx', '.xlsm', '.pptx', '.pptm', '.odt', '.ods', '.odp',
    '.doc', '.xls', '.ppt', '.docm', '.zip'}


class PlainHTML(HTMLParser):
    def __init__(self):
        super().__init__(); self.parts = []; self.hidden = 0
    def handle_starttag(self, tag, attrs):
        if tag in {'script', 'style'}: self.hidden += 1
        if tag in {'p', 'div', 'br', 'tr', 'li', 'h1', 'h2', 'h3'}: self.parts.append('\n')
        if tag in {'td', 'th'}: self.parts.append(' | ')
    def handle_endtag(self, tag):
        if tag in {'script', 'style'}: self.hidden = max(0, self.hidden - 1)
    def handle_data(self, data):
        if not self.hidden: self.parts.append(data)


def plain_html(text):
    parser = PlainHTML(); parser.feed(text)
    return ''.join(parser.parts)


def decode_text(data):
    for encoding in ('utf-8-sig', 'utf-16' if data.startswith((b'\xff\xfe', b'\xfe\xff')) else 'utf-8', 'cp1252'):
        try:
            text = data.decode(encoding)
            if '\x00' not in text: return text
        except UnicodeError: pass
    raise ValueError('This attachment is not readable text. Export a text-based copy.')


def extra_blocks(name, data, depth=0):
    """Return located blocks and explicit extraction limits; never fetch remote content."""
    ext = Path(name).suffix.lower()
    warnings = []
    if ext in {'.mp3', '.wav', '.m4a', '.ogg', '.aac', '.flac'}:
        raise ValueError('Audio needs a text transcript before it can be compared.')
    if ext not in ATTACHMENT_EXTENSIONS and ext not in MEDIA_EXTENSIONS:
        # SAP can download an attachment without its original extension.
        from ingest import extract
        if data.startswith(b'%PDF-'):
            chunks, limits = extract('attachment.pdf', data, depth=depth)
            return [(c.locator, c.text) for c in chunks], limits
        if data.startswith(b'PK\x03\x04'):
            with ZipFile(io.BytesIO(data)) as archive:
                paths=set(archive.namelist())
            inferred='.docx' if 'word/document.xml' in paths else '.xlsx' if 'xl/workbook.xml' in paths else '.pptx' if 'ppt/presentation.xml' in paths else None
            if inferred:
                chunks, limits = extract('attachment'+inferred, data, depth=depth)
                return [(c.locator, c.text) for c in chunks], limits
        try:
            text=data.decode('utf-8-sig')
            if '\x00' not in text and sum(c.isprintable() or c.isspace() for c in text)>=len(text)*.98:
                return [('Text',text)], warnings
        except UnicodeError: pass
    if ext == '.zip':
        if depth>=2: raise ValueError('Nested archives need manual review.')
        blocks=[]
        with ZipFile(io.BytesIO(data)) as archive:
            entries=[entry for entry in archive.infolist() if not entry.is_dir()]
            if len(entries)>100 or sum(entry.file_size for entry in entries)>80*1024*1024:
                raise ValueError('The archive exceeds 100 files or 80 MB expanded.')
            for entry in entries:
                if Path(entry.filename).suffix.lower() in MEDIA_EXTENSIONS: continue
                try:
                    from ingest import extract
                    chunks, limits=extract(entry.filename,archive.read(entry),depth=depth+1)
                    blocks.extend((f'Archive {entry.filename}, {c.locator}',c.text) for c in chunks)
                    warnings.extend(limits)
                except (ValueError,OSError,RuntimeError):
                    warnings.append(f'{name}: archived attachment {entry.filename} could not be read.')
        return blocks,warnings
    if ext in TEXT_EXTENSIONS:
        text = decode_text(data)
        if ext in {'.csv', '.tsv'}:
            return [(f'Row {i}', ' | '.join(row)) for i, row in enumerate(csv.reader(io.StringIO(text), delimiter='\t' if ext == '.tsv' else ','), 1)], warnings
        return [('Text', text)], warnings
    if ext in {'.html', '.htm'}:
        return [('HTML text', plain_html(decode_text(data)))], warnings
    if ext == '.rtf':
        from striprtf.striprtf import rtf_to_text
        return [('RTF text', rtf_to_text(decode_text(data)))], [f'{name}: embedded objects and images were not read.']
    if ext in {'.eml', '.msg'}:
        embedded = []
        if ext == '.eml':
            message = BytesParser(policy=policy.default).parsebytes(data)
            headers = '\n'.join(f'{key}: {message.get(key, "")}' for key in ('From', 'To', 'Date', 'Subject'))
            blocks = [('Email headers', headers)]
            body = message.get_body(preferencelist=('plain', 'html'))
            if body:
                text = body.get_content()
                blocks.append(('Email body', plain_html(text) if body.get_content_subtype() == 'html' else text))
            else: warnings.append(f'{name}: no readable email body was available; verify the original incoming report.')
            for part in message.iter_attachments():
                if part.get_content_maintype() in {'image','video'}: continue
                filename = part.get_filename() or 'embedded-attachment'
                payload = part.get_payload(decode=True)
                if part.get_content_type()=='message/rfc822' and isinstance(part.get_payload(),list):
                    payload=part.get_payload()[0].as_bytes()
                    if Path(filename).suffix.lower()!='.eml': filename+='-forwarded.eml'
                if payload: embedded.append((filename, payload))
        else:
            import extract_msg
            with extract_msg.openMsg(data) as message:
                blocks = [('Email headers', f'From: {message.sender}\nTo: {message.to}\nDate: {message.date}\nSubject: {message.subject}'),
                          ('Email body', message.body or plain_html((message.htmlBody or b'').decode('utf-8', errors='replace')))]
                for attachment in message.attachments:
                    payload = attachment.data
                    if isinstance(payload, bytes): embedded.append((attachment.longFilename or attachment.shortFilename or 'embedded-attachment', payload))
                    elif hasattr(payload,'body'):
                        if depth>=2:
                            warnings.append(f'{name}: nested Outlook message needs manual review.')
                        else:
                            blocks.append(('Embedded Outlook message headers', f'From: {payload.sender}\nDate: {payload.date}\nSubject: {payload.subject}'))
                            blocks.append(('Embedded Outlook message body', payload.body or ''))
                            if payload.attachments: warnings.append(f'{name}: files inside the nested Outlook message need separate text copies.')
                    else:
                        warnings.append(f'{name}: an embedded Outlook item could not be read.')
        for filename, payload in embedded:
            if Path(filename).suffix.lower() in MEDIA_EXTENSIONS: continue
            if depth >= 2:
                warnings.append(f'{name}: nested attachment {filename} needs manual review.'); continue
            try:
                from ingest import extract
                chunks, limits = extract(filename, payload, depth=depth+1)
                blocks.extend((f'Embedded {filename}, {c.locator}', c.text) for c in chunks)
                warnings.extend(limits)
            except (ValueError, OSError):
                warnings.append(f'{name}: embedded attachment {filename} could not be read.')
        return blocks, warnings
    if ext in {'.xlsx', '.xlsm', '.pptx', '.pptm', '.odt', '.ods', '.odp', '.docm'}:
        with ZipFile(io.BytesIO(data)) as archive:
            if sum(e.file_size for e in archive.infolist()) > 80 * 1024 * 1024:
                raise ValueError('The expanded office attachment exceeds 80 MB.')
            if ext in {'.xlsx', '.xlsm'}:
                from openpyxl import load_workbook
                book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
                try:
                    blocks = []
                    for sheet in book:
                        # Ignore stale dimensions; bound actual populated rows below.
                        sheet.reset_dimensions()
                        for i, row in enumerate(sheet.iter_rows(values_only=True), 1):
                            if i > 100000: raise ValueError('The spreadsheet exceeds 100,000 rows.')
                            if len(row)>2048: raise ValueError('The spreadsheet exceeds 2,048 columns.')
                            values = [str(v) if v is not None else '' for v in row]
                            if any(values): blocks.append((f'Sheet {sheet.title}, row {i}', ' | '.join(values)))
                finally: book.close()
                warnings.append(f'{name}: spreadsheet values use saved results; formulas were not recalculated and drawings were not read.')
            else:
                if ext=='.docm':
                    paths=[p for p in archive.namelist() if p.startswith('word/') and p.endswith('.xml') and Path(p).stem.startswith(('document','header','footer','footnotes','endnotes'))]
                else:
                    paths = sorted((p for p in archive.namelist() if (p.startswith('ppt/slides/slide') or p.startswith('ppt/notesSlides/notesSlide')) and p.endswith('.xml')), key=lambda p: (len(p), p)) if ext in {'.pptx', '.pptm'} else ['content.xml']
                blocks = []
                for path in paths:
                    root = ET.fromstring(archive.read(path))
                    text = '\n'.join(''.join(node.itertext()) for node in root.iter() if node.tag.rsplit('}', 1)[-1] in {'p', 'h'})
                    # DrawingML paragraphs have text children, not paragraph text.
                    if text.strip(): blocks.append((path, text))
                warnings.append(f'{name}: images, charts and embedded objects were not interpreted.')
        return blocks, warnings
    if ext in {'.doc', '.xls', '.ppt'}:
        office = shutil.which('libreoffice') or shutil.which('soffice')
        if not office: raise ValueError('This legacy Office attachment needs LibreOffice, or an exported PDF/DOCX/XLSX/PPTX copy.')
        with tempfile.TemporaryDirectory(prefix='gch-convert-') as directory:
            root = Path(directory); source = root / ('input' + ext); source.write_bytes(data)
            result = subprocess.run([office, '-env:UserInstallation='+(root/'profile').as_uri(), '--headless', '--convert-to', 'pdf', '--outdir', str(root), str(source)], capture_output=True, timeout=60)
            output = root / 'input.pdf'
            if result.returncode or not output.exists(): raise ValueError('The Office attachment could not be converted. Export a text-based copy.')
            from ingest import extract
            chunks, limits = extract('converted.pdf', output.read_bytes())
            return [(f'Converted {c.locator}', c.text) for c in chunks], limits
    raise ValueError('This attachment format could not be read. Export a text-based PDF, Office document, email or text copy.')