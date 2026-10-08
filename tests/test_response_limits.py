from copy import deepcopy
import json
from pathlib import Path
import tempfile
import threading
import unittest
from unittest.mock import patch
from config import DEFAULT_MODEL, GROUPS, PROCEDURES, MAX_OUTPUT_TOKENS, MAX_CHECKS_PER_BATCH
from engine import ReviewEngine
from ingest import source_from_bytes
from mdt import MDTClient, ProviderError
from repository import Repository
from responses_fixture import wire_response, invoke_mock, decode_payload
from test_app import MockProvider, PESR
def response(reason='end_turn', text='READY', output_tokens=0):
    return wire_response({'content':[{'type':'text','text':text}], 'model':DEFAULT_MODEL,
            'stop_reason':reason, 'usage':{'input_tokens':10, 'output_tokens':output_tokens}})
class ProviderStopTests(unittest.TestCase):
    def test_length_error_has_a_distinct_reason_and_counts_discarded_usage(self):
        client=MDTClient('fake-key',DEFAULT_MODEL,transport=lambda p:response('max_tokens','PRIVATE_PARTIAL',4096))
        with self.assertRaises(ProviderError) as caught:
            client.message('test',[],max_tokens=4096)
        self.assertEqual(caught.exception.stop_reason,'max_tokens')
        self.assertEqual(caught.exception.output_limit,4096)
        self.assertIn('stop_reason=max_tokens',str(caught.exception))
        self.assertNotIn('PRIVATE_PARTIAL',str(caught.exception))
        self.assertEqual(client.usage['output_tokens'],4096)
    def test_refusal_is_distinct_from_truncation_without_echoing_content(self):
        client=MDTClient('fake-key',DEFAULT_MODEL,transport=lambda p:response('refusal','PRIVATE_EXPLANATION'))
        with self.assertRaises(ProviderError) as caught:
            client.message('test',[])
        self.assertEqual(caught.exception.stop_reason,'refusal')
        self.assertIn('declined this request',str(caught.exception))
        self.assertNotIn('PRIVATE_EXPLANATION',str(caught.exception))
    def test_context_and_pause_stops_are_not_accepted_as_complete(self):
        for reason in ['model_context_window_exceeded','pause_turn']:
            with self.subTest(reason=reason):
                client=MDTClient('fake-key',DEFAULT_MODEL,transport=lambda p:response(reason))
                with self.assertRaises(ProviderError) as caught:client.message('test',[])
                self.assertEqual(caught.exception.stop_reason,'unrecognized')
    def test_unknown_stop_reason_is_not_echoed_or_accepted(self):
        client=MDTClient('fake-key',DEFAULT_MODEL,transport=lambda p:response('PRIVATE_UNKNOWN_VALUE'))
        with self.assertRaises(ProviderError) as caught:client.message('test',[])
        self.assertEqual(caught.exception.stop_reason,'unrecognized')
        self.assertNotIn('PRIVATE_UNKNOWN_VALUE',str(caught.exception))
class RecoveryTests(unittest.TestCase):
    def setUp(self):
        self.engine=ReviewEngine(None)
        self.cancelled=threading.Event()
        self.budget={'remaining':20}
        self.progress=[]
    def invoke(self,client,**kwargs):
        return self.engine._message(client,'test system',[{'role':'user','content':'test input'}],
            self.progress.append,self.cancelled,self.budget,'gathering evidence',max_tokens=4096,**kwargs)
    def test_retry_preserves_original_request_and_drops_partial_content(self):
        calls=[]
        def transport(payload):
            calls.append(deepcopy(payload))
            return response('max_tokens','UNACCEPTED_PARTIAL',4096) if len(calls)==1 else response()
        answer=self.invoke(MDTClient('fake-key',DEFAULT_MODEL,transport=transport),schema={'type':'object'})
        self.assertEqual(answer['stop_reason'],'end_turn')
        self.assertEqual([p['max_output_tokens'] for p in calls],[4096,8192])
        self.assertEqual(calls[0]['input'][1:],calls[1]['input'][1:])
        self.assertEqual(calls[0]['model'],calls[1]['model'])
        self.assertEqual(calls[0]['text'],calls[1]['text'])
        self.assertNotIn('UNACCEPTED_PARTIAL',json.dumps(calls))
        self.assertTrue(self.progress)
    def test_refusal_is_never_retried(self):
        calls=[]
        def transport(payload):calls.append(payload);return response('refusal')
        with self.assertRaisesRegex(ProviderError,'stop_reason=refusal'):
            self.invoke(MDTClient('fake-key',DEFAULT_MODEL,transport=transport))
        self.assertEqual(len(calls),1)
    def test_retries_stop_at_the_output_cap(self):
        calls=[]
        def transport(payload):calls.append(payload);return response('max_tokens')
        with self.assertRaisesRegex(ProviderError,'retries are exhausted') as caught:
            self.invoke(MDTClient('fake-key',DEFAULT_MODEL,transport=transport))
        self.assertEqual([p['max_output_tokens'] for p in calls],[4096,8192,MAX_OUTPUT_TOKENS])
        self.assertEqual(caught.exception.stop_reason,'max_tokens')
        self.assertIn('gathering evidence',str(caught.exception))
    def test_cancellation_prevents_the_next_retry(self):
        calls=[]
        def transport(payload):
            calls.append(payload);self.cancelled.set();return response('max_tokens')
        with self.assertRaises(InterruptedError):
            self.invoke(MDTClient('fake-key',DEFAULT_MODEL,transport=transport))
        self.assertEqual(len(calls),1)
    def test_request_budget_stops_before_another_provider_call(self):
        self.budget['remaining']=0
        with self.assertRaisesRegex(ProviderError,'request limit'):
            self.invoke(MDTClient('fake-key',DEFAULT_MODEL,transport=lambda p:self.fail('No budget remains')))
    def test_http_parameter_rejection_does_not_trigger_length_retries(self):
        calls=[]
        class Client:
            def message(self,*args,**kwargs):
                calls.append(kwargs)
                raise ProviderError('MDT rejected max_tokens',status=400,field='max_tokens')
        with self.assertRaisesRegex(ProviderError,'MDT rejected'):
            self.invoke(Client())
        self.assertEqual(len(calls),1)
    def test_503_retries_the_same_turn_after_30_seconds(self):
        from test_gateway import rejected
        calls=[]
        def transport(payload):
            calls.append(deepcopy(payload))
            if len(calls)==1:raise rejected('Service unavailable',503)
            return response()
        with patch.object(self.cancelled,'wait',return_value=False) as waiting:
            self.invoke(MDTClient('fake-key',transport=transport))
        waiting.assert_called_once_with(30)
        self.assertEqual(calls[0],calls[1])
        self.assertEqual(self.budget['remaining'],18)
    def test_excess_native_and_portable_calls_retry_without_replaying_rejected_calls(self):
        from test_gateway import TOOLS
        for portable in [False,True]:
            with self.subTest(portable=portable):
                calls=[]
                def transport(payload):
                    calls.append(deepcopy(payload))
                    if len(calls)>1:
                        return response(text='{"tool_calls":[],"text":"READY"}' if portable else 'READY')
                    if portable:
                        return response(text=json.dumps({'tool_calls':[{'name':'read_section','arguments':{'source_id':'s1'}}]*9,'text':''}))
                    return wire_response({'content':[{'type':'tool_use','id':'discarded-'+str(i),'name':'read_section','input':{'source_id':'s1'}} for i in range(9)]})
                client=MDTClient('fake-key',transport=transport)
                if portable:client.tool_mode='validated_json'
                with patch.object(self.cancelled,'wait',return_value=False) as waiting:
                    self.invoke(client,tools=TOOLS)
                waiting.assert_called_once_with(30)
                self.assertEqual(calls[0]['input'][1:],calls[1]['input'][1:])
                self.assertIn('at most 8',calls[1]['input'][0]['content'])
                self.assertNotIn('discarded-',json.dumps(calls))
    def test_persistent_503_stops_after_three_retries(self):
        from test_gateway import rejected
        calls=[]
        def transport(payload):calls.append(payload);raise rejected('Unavailable',503)
        with patch.object(self.cancelled,'wait',return_value=False) as waiting:
            with self.assertRaises(ProviderError):self.invoke(MDTClient('fake-key',transport=transport))
        self.assertEqual(len(calls),4)
        self.assertEqual(waiting.call_count,3)
    def test_cancellation_during_recovery_wait_stops_without_another_request(self):
        from test_gateway import rejected
        calls=[]
        def transport(payload):calls.append(payload);raise rejected('Unavailable',503)
        with patch.object(self.cancelled,'wait',return_value=True):
            with self.assertRaises(InterruptedError):self.invoke(MDTClient('fake-key',transport=transport))
        self.assertEqual(len(calls),1)
class ReviewRecoveryTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory()
        self.repo=Repository(Path(self.temp.name))
        for docid,_ in PROCEDURES:
            content=f'{docid} revision TEST ONLY\nUse the supplied event identifier for this synthetic exercise. No real procedure is represented.\n'.encode()
            source,warnings=source_from_bytes(docid+'.txt',content,docid,'procedure',docid,revision='TEST',approved=True)
            self.repo.put(source,warnings)
        self.connection={'token':'fake-key','model':DEFAULT_MODEL,'auth_style':'bearer'}
    def tearDown(self):
        self.repo.db.close();self.temp.cleanup()
    def run_review(self,factory,job_id='test-review'):
        return ReviewEngine(self.repo,factory).run(job_id,'TRAIN-001','Intake','Synthetic test only',
            [('PESR','PESR.txt',PESR)],self.connection,lambda *args:None,threading.Event())
    def test_503_in_a_later_packet_keeps_completed_findings_and_current_conversation(self):
        from test_gateway import rejected
        calls=[];failed=False;failed_index=None
        def factory(token,model,style,**kwargs):
            mock=MockProvider(token,model,style)
            def transport(payload):
                nonlocal failed,failed_index
                calls.append(deepcopy(payload))
                task=json.loads(payload['input'][1]['content'])
                if not failed and task['expected_checks'][0]['id']=='gfe' and 'text' in payload:
                    failed=True;failed_index=len(calls)-1;raise rejected('Service unavailable',503)
                return invoke_mock(mock,payload)
            return MDTClient(token,model,style,transport=transport,**kwargs)
        with patch('threading.Event.wait',return_value=False) as waiting:
            result=self.run_review(factory,'retry-later-packet')
        self.assertTrue(failed)
        waiting.assert_called_once_with(30)
        self.assertEqual(calls[failed_index],calls[failed_index+1])
        self.assertEqual(len(result['review']['findings']),13)
        first=[p for p in calls if json.loads(p['input'][1]['content'])['expected_checks'][0]['id']=='intake' and 'text' in p]
        self.assertEqual(len(first),2)
        self.assertTrue((self.repo.directory/'reports'/'retry-later-packet'/'review.docx').exists())
    def test_full_review_recovers_a_cut_off_tool_or_final_turn_and_writes_all_checks(self):
        for phase in ['retrieval','final']:
            with self.subTest(phase=phase):
                calls=[];truncated=False
                def factory(token,model,style,**kwargs):
                    mock=MockProvider(token,model,style)
                    def transport(payload):
                        nonlocal truncated
                        calls.append(deepcopy(payload))
                        is_final='text' in payload
                        if not truncated and (is_final==(phase=='final')):
                            truncated=True
                            answer=response('max_tokens','UNACCEPTED_PARTIAL',payload['max_output_tokens'])
                            answer['output'].append({'type':'function_call','call_id':'unaccepted-tool','name':'read_section','arguments':'{}'})
                            return answer
                        return invoke_mock(mock,payload)
                    return MDTClient(token,model,style,transport=transport,**kwargs)
                result=self.run_review(factory,'review-'+phase)
                ids=[f['check_id'] for f in result['review']['findings']]
                self.assertEqual(set(ids),{c[0] for g in GROUPS for c in g['checks']})
                self.assertEqual(len(ids),len(set(ids)))
                self.assertTrue((self.repo.directory/'reports'/('review-'+phase)/'review.docx').read_bytes().startswith(b'PK'))
                self.assertTrue(truncated)
                self.assertNotIn('unaccepted-tool',json.dumps(calls))
                self.assertNotIn('UNACCEPTED_PARTIAL',json.dumps(calls))
                self.assertGreater(result['usage']['output_tokens'],0)
                for call in calls:
                    task=json.loads(call['input'][1:][0]['content'])
                    self.assertLessEqual(len(task['expected_checks']),MAX_CHECKS_PER_BATCH)
                    if 'text' not in call:self.assertIn('CURRENT PHASE: gather evidence',call['input'][0]['content'])
    def test_complete_review_with_documented_text_only_gateway(self):
        from test_gateway import rejected
        from responses_fixture import response as wire_text
        clients=[];requests=[]
        def factory(token,model,style,**kwargs):
            mock=MockProvider(token,model,style)
            def transport(payload):
                requests.append(payload)
                if 'tools' in payload:raise rejected('tools is not supported')
                if 'text' in payload:raise rejected("Unknown parameter: 'text'")
                system,messages=decode_payload(payload)
                for msg in messages:
                    if isinstance(msg['content'],str) and msg['content'].startswith('['):
                        msg['content']=json.loads(msg['content'])
                final='Return only a JSON object matching this schema' in system
                answer=mock.message(system,messages,tools=[{'name':'read'}],schema=True if final else None)
                if final:return wire_response(answer)
                plan={'tool_calls':[{'name':c['name'],'arguments':c['input']} for c in answer['content'] if c['type']=='tool_use'],
                      'text':'READY' if answer['stop_reason']=='end_turn' else ''}
                return wire_text(json.dumps(plan))
            client=MDTClient(token,model,style,transport=transport,**kwargs);clients.append(client);return client
        result=self.run_review(factory,'portable-review')
        self.assertEqual(len(result['review']['findings']),13)
        self.assertEqual(clients[0].tool_mode,'validated_json')
        self.assertEqual(clients[0].structured_mode,'validated_json')
        self.assertTrue(all(p['model']==DEFAULT_MODEL and p['store'] is False for p in requests))
        self.assertTrue((self.repo.directory/'reports'/'portable-review'/'review.docx').read_bytes().startswith(b'PK'))
        self.assertTrue(any('JSON plan' in text for text in result['review']['limitations']))
    def test_gateway_reported_model_names_remain_visible_in_report_evidence(self):
        requests=[]
        def factory(token,model,style,**kwargs):
            mock=MockProvider(token,model,style)
            def transport(payload):
                requests.append(payload)
                answer=invoke_mock(mock,payload)
                answer['model']='internal-deployment-snapshot'
                return answer
            return MDTClient(token,model,style,transport=transport,**kwargs)
        result=self.run_review(factory,'alias-review')
        self.assertEqual(len(result['review']['findings']),13)
        out=self.repo.directory/'reports'/'alias-review'
        evidence=json.loads((out/'evidence.json').read_text())
        self.assertIn('Requested: '+DEFAULT_MODEL,evidence['model_version'])
        self.assertIn('gateway reported: internal-deployment-snapshot',evidence['model_version'])
        usage=json.loads((out/'context_usage.json').read_text())
        self.assertTrue(all(x['requested_model']==DEFAULT_MODEL and x['reported_model']=='internal-deployment-snapshot' for x in usage))
        self.assertEqual({p['model'] for p in requests},{DEFAULT_MODEL})
        self.assertTrue((out/'review.docx').read_bytes().startswith(b'PK'))
    def test_refusal_creates_no_word_document_and_makes_one_call(self):
        calls=[]
        def factory(token,model,style,**kwargs):
            def transport(payload):calls.append(payload);return response('refusal')
            return MDTClient(token,model,style,transport=transport,**kwargs)
        with self.assertRaisesRegex(ProviderError,'stop_reason=refusal'):self.run_review(factory)
        self.assertEqual(len(calls),1)
        self.assertFalse((self.repo.directory/'reports'/'test-review'/'review.docx').exists())
    def test_specific_validation_feedback_repairs_response_before_word_output(self):
        observed=[];damaged=False
        def factory(token,model,style,**kwargs):
            self.assertEqual(kwargs['environment'],'production')
            mock=MockProvider(token,model,style)
            def transport(payload):
                nonlocal damaged
                answer=invoke_mock(mock,payload)
                if 'text' in payload:
                    last=payload['input'][1:][-1]['content']
                    if isinstance(last,str) and last.startswith('{'):
                        correction=json.loads(last)
                        observed.extend(correction.get('validation_errors',[]))
                    if not damaged:
                        damaged=True
                        data=json.loads(answer['output'][0]['content'][0]['text'])
                        data['findings'][0]['feedback']=['PRIVATE_RESPONSE_TEXT']
                        answer['output'][0]['content'][0]['text']=json.dumps(data)
                return answer
            return MDTClient(token,model,style,transport=transport,**kwargs)
        self.connection['environment']='development'  # Stale preference must not reach the provider.
        result=self.run_review(factory)
        self.assertEqual(observed[0]['code'],'string_type')
        self.assertEqual(observed[0]['path'],'findings.0.feedback')
        self.assertNotIn('PRIVATE_RESPONSE_TEXT',json.dumps(observed))
        self.assertEqual(len(result['review']['findings']),13)
        for finding in result['review']['findings']:
            for cite in finding['record_evidence']:
                self.assertIn(cite['quote'],PESR.decode())
        self.assertTrue((self.repo.directory/'reports'/'test-review'/'review.docx').read_bytes().startswith(b'PK'))
    def test_bad_citations_after_a_length_retry_still_block_the_word_document(self):
        truncated=False;final_calls=0
        def factory(token,model,style,**kwargs):
            mock=MockProvider(token,model,style)
            def transport(payload):
                nonlocal truncated,final_calls
                if not truncated:
                    truncated=True;return response('max_tokens')
                answer=invoke_mock(mock,payload)
                if 'text' in payload:
                    final_calls+=1
                    data=json.loads(answer['output'][0]['content'][0]['text'])
                    data['findings'][0]['record_evidence'][0]='FABRICATED_EVIDENCE_ID'
                    answer['output'][0]['content'][0]['text']=json.dumps(data)
                return answer
            return MDTClient(token,model,style,transport=transport,**kwargs)
        with self.assertRaisesRegex(ProviderError,'unknown_evidence_id'):self.run_review(factory)
        self.assertEqual(final_calls,2)
        self.assertFalse((self.repo.directory/'reports'/'test-review'/'review.docx').exists())
    def test_long_gfe_feedback_completes_without_format_retries(self):
        text=('Synthetic GFE explanation with conditions and uncertainty retained. '*30)+'FINAL CONDITION REMAINS.'
        gfe_calls=0
        def factory(token,model,style,**kwargs):
            mock=MockProvider(token,model,style)
            def transport(payload):
                nonlocal gfe_calls
                answer=invoke_mock(mock,payload)
                if 'text' in payload:
                    data=json.loads(answer['output'][0]['content'][0]['text'])
                    for finding in data['findings']:
                        if finding['check_id']=='gfe':
                            gfe_calls+=1
                            finding['feedback']=text
                    answer['output'][0]['content'][0]['text']=json.dumps(data)
                return answer
            return MDTClient(token,model,style,transport=transport,**kwargs)
        result=self.run_review(factory)
        self.assertEqual(gfe_calls,2)
        self.assertEqual(len(result['review']['findings']),13)
        self.assertEqual(next(f['feedback'] for f in result['review']['findings'] if f['check_id']=='gfe'),text)
        saved=json.loads((self.repo.directory/'reports'/'test-review'/'review.json').read_text())
        self.assertEqual(next(f['feedback'] for f in saved['findings'] if f['check_id']=='gfe'),text)
        from docx import Document
        document=Document(self.repo.directory/'reports'/'test-review'/'review.docx')
        contents=[p.text for p in document.paragraphs]
        contents.extend(cell.text for table in document.tables for row in table.rows for cell in row.cells)
        self.assertIn(text,contents)
if __name__=='__main__':unittest.main()