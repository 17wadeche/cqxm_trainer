"""Synthetic tests only. No MDT-GPT or GCH network calls."""
import base64
import io
import json
from pathlib import Path
import tempfile
import threading
import time
import unittest
from unittest.mock import patch
import urllib.error
import urllib.request

from docx import Document
from pypdf import PdfWriter

from app import Application, make_server
from config import ENDPOINT, DEFAULT_MODEL, PROCEDURES, GROUPS
from engine import ReviewEngine
from ingest import extract, source_from_bytes
from responses_fixture import response, wire_response
from mdt import MDTClient, ProviderError
from models import EvidenceManifest, Review, validate_traceability
from repository import Repository

PESR = b'Product Event Summary Report\nEvent ID: TRAIN-001\nSYNTHETIC TEST ONLY. Intake is in progress. No real patient or procedure content.\n'


def encoded(name, content, kind=None):
    return {'name': name, 'data': base64.b64encode(content).decode(), **({'kind': kind} if kind else {})}


class MockProvider:
    """Imitates internal tool_use -> tool_result -> structured final output."""
    payloads = []
    def __init__(self, token, model, auth_style, **kwargs):
        self.model = model
        self.usage = {'input_tokens': 100, 'output_tokens': 50}

    def message(self, system, messages, tools=None, schema=None, max_tokens=1, diagnostic=False):
        self.payloads.append({'tools': tools, 'schema': schema, 'messages': messages})
        if 'connection test' in system.lower():
            return {'content':[{'type':'text','text':'{"status":"ok"}'}], 'stop_reason':'end_turn', 'model':self.model}
        task = json.loads(messages[0]['content'])
        if tools and not schema:
            has_result=any(isinstance(m['content'],list) and any(c.get('type')=='tool_result' for c in m['content']) for m in messages)
            if not has_result:
                return {'content':[{'type':'tool_use','id':'t1','name':'search_procedures','input':{'query':'synthetic identifier'}}], 'stop_reason':'tool_use'}
            return {'content':[{'type':'text','text':'Evidence retrieved.'}], 'stop_reason':'end_turn'}
        record = next(s for s in task['sources'] if s['kind']=='record')
        findings=[]
        for check in task['expected_checks']:
            rule=next(s for s in task['sources'] if s['document_id']==check['required_documents'][0])
            findings.append({'check_id':check['id'],'record_element':check['focus'],
                'what_was_done':'Synthetic test record supplied; no clinical interpretation is performed.',
                'record_evidence':[record['start']['excerpts'][0]['evidence_id']],
                'procedure_references':[rule['start']['excerpts'][0]['evidence_id']],
                'assessment':'needs_trainer_review','feedback':'Synthetic test output. Verify source interpretation with a trainer.',
                'priority':'low','study_action':'Review the synthetic example.'})
            if check['id']=='consistency':
                mdr=next(s for s in task['sources'] if s['document_id']=='MDR')
                findings[-1]['record_evidence'].append(mdr['start']['excerpts'][0]['evidence_id'])
        text=json.dumps({'overall_summary':'Synthetic mock response.','limitations':['Mock provider; no actual AI review.'],'findings':findings})
        return {'content':[{'type':'text','text':text}], 'stop_reason':'end_turn'}


class AppTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory()
        self.directory=Path(self.temp.name)
        self.app=Application(self.directory,port=0,token='test-local-session',provider_factory=MockProvider,bundled_directory=None)
        self.server=make_server(self.app)
        self.thread=threading.Thread(target=self.server.serve_forever,daemon=True)
        self.thread.start()
        self.base=f'http://127.0.0.1:{self.app.port}'

    def tearDown(self):
        self.server.shutdown();self.server.server_close();self.app.pool.shutdown(wait=True)
        self.app.repo.db.close();self.temp.cleanup()

    def request(self, path, data=None, token='test-local-session', origin=None):
        headers={'Authorization':'Bearer '+token}
        if origin:headers['Origin']=origin
        if data is not None:headers['Content-Type']='application/json'
        req=urllib.request.Request(self.base+path, data=json.dumps(data).encode() if data is not None else None, headers=headers)
        try:
            response=urllib.request.urlopen(req,timeout=10)
        except urllib.error.HTTPError as exc:
            return exc.code, json.loads(exc.read())
        with response:
            body=response.read()
            if response.headers.get('Content-Type','').startswith('application/json'):
                body=json.loads(body)
            return response.status,body

    def configure(self):
        code,_=self.request('/api/connection',{'token':'synthetic-mdt-secret','model':DEFAULT_MODEL,'auth_style':'bearer'})
        self.assertEqual(code,200)
        for docid,_ in PROCEDURES[:6]:
            content=f'{docid} revision TEST ONLY\nUse the supplied event identifier for this synthetic exercise. No real Medtronic procedure is represented.\n'.encode()
            code,body=self.request('/api/procedures',{'document_id':docid,'revision':'TEST','approved':True,'file':encoded(docid+'.txt',content)})
            self.assertEqual(code,200,body)

    def wait_job(self, job):
        deadline=time.monotonic()+12
        while time.monotonic()<deadline:
            _,current=self.request('/api/jobs/'+job['id'])
            if current['status'] not in {'queued','running'}:
                return current
            time.sleep(.02)
        self.fail('Review did not finish')

    def test_api_requires_local_token(self):
        code,_=self.request('/api/status',token='wrong')
        self.assertEqual(code,403)

    def test_manual_connection_cannot_select_development(self):
        code,result=self.request('/api/connection',{'model':DEFAULT_MODEL,'token':'fake-key','environment':'development'})
        self.assertEqual(code,400)
        self.assertIn('only connects to MDT-GPT Production',json.dumps(result))
        self.assertEqual(self.app.status()['environment'],'production')

    def test_untrusted_origin_is_rejected_even_with_token(self):
        code,_=self.request('/api/status',origin='https://example.com')
        self.assertEqual(code,403)

    def test_demo_download_is_valid_word(self):
        code,job=self.request('/api/demo',{})
        self.assertEqual(code,200)
        self.assertTrue(job['result']['demo'])
        code,word=self.request('/api/jobs/'+job['id']+'/word')
        self.assertEqual(code,200)
        self.assertTrue(word.startswith(b'PK'))
        doc=Document(io.BytesIO(word))
        self.assertTrue(any('SYNTHETIC DEMO' in p.text for p in doc.paragraphs))

    def test_api_token_is_not_persisted_or_returned(self):
        self.request('/api/connection',{'token':'secret-that-must-not-be-written','model':DEFAULT_MODEL})
        saved=(self.directory/'settings.json').read_text()
        _,status=self.request('/api/status')
        self.assertNotIn('secret-that-must-not-be-written',saved+json.dumps(status))

    def test_procedure_requires_approval(self):
        code,_=self.request('/api/procedures',{'document_id':'027-P043','revision':'TEST','approved':False,'file':encoded('p.txt',b'027-P043 synthetic text for testing only and no policy rules.')})
        self.assertEqual(code,400)

    def test_wrong_procedure_identity_rejected(self):
        code,_=self.request('/api/procedures',{'document_id':'027-P043','revision':'TEST','approved':True,'file':encoded('p.txt',b'Different procedure document. This is synthetic text for testing only.')})
        self.assertEqual(code,400)

    def test_real_flow_uses_tools_and_produces_cited_report(self):
        MockProvider.payloads=[]
        self.configure()
        code,job=self.request('/api/reviews',{'record_id':'TRAIN-001','stage':'Intake in progress','policy_selection':'Synthetic test revisions','files':[encoded('PESR.txt',PESR,'PESR')]})
        self.assertEqual(code,202,job)
        done=self.wait_job(job)
        self.assertEqual(done['status'],'ready',done)
        findings=done['result']['review']['findings']
        self.assertEqual(len(findings),13)
        for key in ['consistency']:
            self.assertEqual(next(f for f in findings if f['check_id']==key)['assessment'],'not_assessable')
        evidence=EvidenceManifest.model_validate_json((self.directory/'reports'/job['id']/'evidence.json').read_text())
        validate_traceability(Review.model_validate(done['result']['review']),evidence)
        self.assertTrue(any(p['tools'] for p in MockProvider.payloads))
        self.assertTrue(any(p['schema'] for p in MockProvider.payloads))
        _,word=self.request('/api/jobs/'+job['id']+'/word')
        self.assertTrue(word.startswith(b'PK'))

    def test_wrong_record_stops_before_model(self):
        self.configure()
        MockProvider.payloads=[]
        _,job=self.request('/api/reviews',{'record_id':'TRAIN-002','stage':'Intake','policy_selection':'Test only','files':[encoded('PESR.txt',PESR,'PESR')]})
        done=self.wait_job(job)
        self.assertEqual(done['status'],'failed')
        self.assertEqual(MockProvider.payloads,[])

    def test_capture_cannot_read_outside_downloads(self):
        self.configure()
        self.app.downloads=self.directory/'Downloads';self.app.downloads.mkdir()
        outside=self.directory/'outside.txt';outside.write_bytes(PESR)
        _,cap=self.request('/api/captures',{'record_id':'TRAIN-001','stage':'Intake'})
        code,body=self.request('/api/captures/'+cap['capture_id']+'/complete',{'path':str(outside)})
        self.assertEqual(code,400)
        self.assertIn('outside',body['error'])

    def test_capture_fresh_report_finishes_full_flow(self):
        self.configure()
        self.app.downloads=self.directory/'Downloads';self.app.downloads.mkdir()
        _,cap=self.request('/api/captures',{'record_id':'TRAIN-001','stage':'Intake'})
        report=self.app.downloads/'fresh.txt';report.write_bytes(PESR)
        code,job=self.request('/api/captures/'+cap['capture_id']+'/complete',{'path':str(report)})
        self.assertEqual(code,202,job)
        self.assertEqual(self.wait_job(job)['status'],'ready')
        code,_=self.request('/api/captures/'+cap['capture_id']+'/complete',{'path':str(report)})
        self.assertEqual(code,400)

    def test_malformed_file_object_returns_error(self):
        code,_=self.request('/api/reviews',{'files':[None]})
        self.assertEqual(code,400)


class ExtractionAndProviderTests(unittest.TestCase):
    def test_nested_word_table_is_extracted(self):
        doc=Document();doc.add_paragraph('Synthetic extraction test only.')
        outer=doc.add_table(rows=1,cols=1)
        inner=outer.cell(0,0).add_table(rows=1,cols=1);inner.cell(0,0).text='Nested evidence survives.'
        buffer=io.BytesIO();doc.save(buffer)
        chunks,_=extract('nested.docx',buffer.getvalue())
        self.assertTrue(any('Nested evidence survives.' in c.text and 'table' in c.locator for c in chunks))

    def test_scanned_or_blank_pdf_is_rejected(self):
        writer=PdfWriter();writer.add_blank_page(width=612,height=792);out=io.BytesIO();writer.write(out)
        with self.assertRaisesRegex(ValueError,'No usable text'):
            extract('blank.pdf',out.getvalue())

    def test_wrong_report_type_rejected(self):
        with self.assertRaisesRegex(ValueError,'not identifiable'):
            source_from_bytes('other.txt',b'Event ID: TRAIN-001. This is a random document and not the expected report.', 'x','record','PESR','TRAIN-001')

    def test_adapter_uses_provider_endpoint_and_bearer_header(self):
        class Response:
            def __enter__(self):return self
            def __exit__(self,*args):pass
            def read(self,*args):return json.dumps(response('ok')).encode()
        seen=[]
        class Opener:
            def open(self,req,timeout):seen.append(req);return Response()
        with patch('urllib.request.build_opener',return_value=Opener()):
            MDTClient('test-secret',DEFAULT_MODEL,'bearer').message('Test',[{'role':'user','content':'hi'}],tools=[{'name':'example','input_schema':{'type':'object'}}],schema={'type':'object'})
        self.assertEqual(seen[0].full_url,ENDPOINT)
        self.assertEqual(seen[0].get_header('Authorization'),'Bearer test-secret')
        body=json.loads(seen[0].data)
        self.assertFalse(body['store'])
        self.assertEqual(body['text']['format']['type'],'json_schema')
        self.assertEqual(body['tool_choice'], 'none')

    def test_truncated_model_response_rejected(self):
        client=MDTClient('x',DEFAULT_MODEL,transport=lambda p:response(stop_reason='max_tokens'))
        with self.assertRaisesRegex(ProviderError,'incomplete'):
            client.message('test',[{'role':'user','content':'hi'}])

    def test_gateway_deployment_metadata_does_not_change_request_model(self):
        client=MDTClient('x',DEFAULT_MODEL,transport=lambda p:response(model='some-other-model'))
        answer=client.message('test',[{'role':'user','content':'hi'}])
        self.assertEqual(answer['reported_model'],'some-other-model')
        self.assertEqual(answer['requested_model'],DEFAULT_MODEL)
        self.assertEqual(client.model,DEFAULT_MODEL)

    def test_context_guard_runs_before_transport(self):
        client=MDTClient('x',DEFAULT_MODEL,transport=lambda p:self.fail('Should not call transport'))
        with self.assertRaisesRegex(ProviderError,'context budget'):
            client.message('test',[{'role':'user','content':'x'*240000}])

    def test_unknown_tool_cannot_execute_code(self):
        with tempfile.TemporaryDirectory() as tmp:
            repo=Repository(Path(tmp));engine=ReviewEngine(repo)
            with self.assertRaisesRegex(ValueError,'Unknown tool'):
                engine._tools('execute_python',{'code':'print(1)'},[])
            repo.db.close()

    def test_read_section_rejects_outside_source(self):
        with tempfile.TemporaryDirectory() as tmp:
            repo=Repository(Path(tmp));engine=ReviewEngine(repo)
            with self.assertRaisesRegex(ValueError,'outside'):
                engine._tools('read_section',{'source_id':'secret','chunk_id':'c1'},[])
            repo.db.close()


if __name__=='__main__':unittest.main()
