"""Word presentation must preserve prose that exceeds table-sized targets."""
from copy import deepcopy
import json
from pathlib import Path
import tempfile
import unittest

from docx import Document
from pydantic import ValidationError

from config import ROOT
from models import EvidenceManifest, Review
from render_report import render


class LongFeedbackTests(unittest.TestCase):
    def setUp(self):
        self.data=json.loads((ROOT/'examples/demo_review.json').read_text())
        self.evidence=EvidenceManifest.model_validate_json((ROOT/'examples/demo_evidence.json').read_text())

    def make_doc(self,data):
        review=Review.model_validate(data)
        with tempfile.TemporaryDirectory() as directory:
            path=Path(directory)/'review.docx'
            render(review,self.evidence,path,detailed=True)
            return Document(path)

    def test_long_feedback_moves_to_detail_section_without_losing_its_ending(self):
        text=('Synthetic long explanation. Preserve every condition and uncertainty. '*32)+' FINAL CONDITION REMAINS.'
        self.data['findings'][0]['feedback']=text
        doc=self.make_doc(self.data)
        body='\n'.join(p.text for p in doc.paragraphs)
        self.assertIn('Detailed review feedback',body)
        self.assertIn(text,body)
        self.assertEqual(doc.tables[1].cell(1,5).text,'See item 1 in Detailed review feedback.')
        self.assertEqual(self.data['findings'][0]['feedback'],text)

    def test_other_long_narrative_fields_also_preserve_full_text(self):
        first=self.data['findings'][0]
        first.update(record_element='Synthetic review area '*10,what_was_done='Synthetic observation. '*35,
                     study_action='Synthetic learning activity. '*35,priority='low')
        doc=self.make_doc(self.data)
        body='\n'.join(p.text for p in doc.paragraphs)
        for field in ['record_element','what_was_done','study_action']:
            self.assertIn(first[field].strip(),body)
        self.assertIn('See item 1',doc.tables[1].cell(1,1).text)
        self.assertIn('See item 1',doc.tables[1].cell(1,2).text)
        self.assertTrue(any('See item 1' in row.cells[2].text for row in doc.tables[2].rows))

    def test_short_feedback_stays_in_existing_table_layout(self):
        doc=self.make_doc(self.data)
        self.assertNotIn('Detailed review feedback',[p.text for p in doc.paragraphs])
        self.assertEqual(doc.tables[1].cell(1,5).text,self.data['findings'][0]['feedback'])

    def test_many_short_lines_also_use_flowing_details(self):
        text='\n'.join('Synthetic line '+str(i) for i in range(10))
        self.assertLess(len(text),600)
        self.data['findings'][0]['feedback']=text
        doc=self.make_doc(self.data)
        self.assertIn(text,[p.text for p in doc.paragraphs])
        self.assertIn('See item 1',doc.tables[1].cell(1,5).text)

    def test_long_feedback_cannot_bypass_citation_validation(self):
        self.data['findings'][0]['feedback']='Synthetic detail. '*200
        self.data['findings'][0]['record_evidence'][0]['quote']='A fabricated quote.'
        with self.assertRaisesRegex(ValueError,'Unverified excerpt'):
            self.make_doc(self.data)

    def test_empty_or_nontext_feedback_still_fails_validation(self):
        for value in ['', '  ', None, [], 3]:
            with self.subTest(value=value):
                data=deepcopy(self.data);data['findings'][0]['feedback']=value
                with self.assertRaises(ValidationError):Review.model_validate(data)


if __name__=='__main__':unittest.main()
