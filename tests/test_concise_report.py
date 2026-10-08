from copy import deepcopy
import json
from pathlib import Path
import tempfile
import unittest
from docx import Document
from config import ROOT
from models import EvidenceManifest, Review
from render_report import compact_references, render
class ConciseReportTests(unittest.TestCase):
    def setUp(self):
        self.data=json.loads((ROOT/'examples/demo_review.json').read_text())
        self.evidence=EvidenceManifest.model_validate_json((ROOT/'examples/demo_evidence.json').read_text())
    def make_doc(self, detailed=False):
        with tempfile.TemporaryDirectory() as directory:
            path=Path(directory)/'review.docx'
            render(Review.model_validate(self.data),self.evidence,path,detailed=detailed)
            return Document(path)
    def test_default_keeps_all_assessments_without_source_appendices(self):
        original=deepcopy(self.data)
        doc=self.make_doc()
        body='\n'.join(p.text for p in doc.paragraphs)
        self.assertEqual(len(doc.tables[0].rows),len(self.data['findings'])+1)
        for i,finding in enumerate(self.data['findings'],1):
            done=finding['assessment']=='done_properly'
            self.assertEqual(doc.tables[0].cell(i,2).text,'' if done else finding['what_was_done'])
            self.assertEqual(doc.tables[0].cell(i,5).text,'' if done else finding['feedback'])
        for heading in ['Evidence and trainer notes','Detailed review feedback','Scope and limitations']:
            self.assertNotIn(heading,body)
        self.assertNotIn(self.data['findings'][0]['procedure_references'][0]['quote'],body)
        self.assertIn('matching MDR was not supplied',body)
        self.assertEqual(self.data,original)
    def test_done_properly_is_blank_with_green_check_and_no_study_priority(self):
        self.data['findings'][0].update(assessment='done_properly',priority='low',study_action='Do not repeat this action')
        original=deepcopy(self.data)
        doc=self.make_doc()
        self.assertEqual(doc.tables[0].cell(1,2).text,'')
        self.assertEqual(doc.tables[0].cell(1,5).text,'')
        assessment=doc.tables[0].cell(1,4)
        self.assertEqual(assessment.text,'✓ Done properly')
        self.assertEqual(str(assessment.paragraphs[0].runs[0].font.color.rgb),'008000')
        self.assertNotIn('Do not repeat this action','\n'.join(c.text for t in doc.tables for row in t.rows for c in row.cells))
        self.assertEqual(self.data,original)
    def test_thirteen_findings_remain_and_only_top_five_actions_repeat(self):
        template=deepcopy(self.data['findings'][0])
        findings=[]
        priorities=['low','medium','high','high','medium','high','low','high','none','high','medium','low','high']
        for i,priority in enumerate(priorities):
            f=deepcopy(template)
            f.update(check_id=f'check{i}',record_element=f'Area {i}',assessment='needs_trainer_review',priority=priority,study_action=f'Action {i}')
            findings.append(f)
        self.data['findings']=findings
        self.evidence.expected_check_ids=[f['check_id'] for f in findings]
        doc=self.make_doc()
        self.assertEqual(len(doc.tables[0].rows),14)
        self.assertEqual(len(doc.tables[1].rows),6)
        self.assertEqual([r.cells[2].text for r in doc.tables[1].rows[1:]],
                         ['Action 2','Action 3','Action 5','Action 7','Action 9'])
    def test_compact_references_deduplicate_chunks_without_losing_pages(self):
        sources={s.id:s for s in self.evidence.sources}
        rule=sources['rule-demo']
        first=rule.chunks[0]
        first.locator='PDF page 2, part 1'
        rule.chunks.extend([first.model_copy(update={'id':'part2','locator':'PDF page 2, part 2'}),
                            first.model_copy(update={'id':'page9','locator':'PDF page 9'})])
        finding=Review.model_validate(self.data).findings[0]
        ref=finding.procedure_references[0]
        finding.procedure_references.extend([ref.model_copy(update={'chunk_id':'part2'}),
                                            ref.model_copy(update={'chunk_id':'page9'})])
        self.assertEqual(compact_references(finding,sources),'DEMO-RULE Rev DEMO ONLY\npp. 2, 9')
    def test_long_feedback_is_not_cut_off_by_concise_layout(self):
        text=('Important condition with uncertainty. '*200)+'FINAL CONDITION.'
        self.data['findings'][0]['feedback']=text
        self.data['findings'][0]['assessment']='needs_trainer_review'
        doc=self.make_doc()
        self.assertEqual(doc.tables[0].cell(1,5).text,text)
    def test_detailed_option_retains_full_evidence_and_limitations(self):
        doc=self.make_doc(detailed=True)
        body='\n'.join(p.text for p in doc.paragraphs)
        self.assertIn('Evidence and trainer notes',body)
        self.assertIn(self.data['findings'][0]['procedure_references'][0]['quote'],body)
        self.assertIn(self.evidence.limitations[0],body)
    def test_bad_citation_still_blocks_concise_report(self):
        self.data['findings'][0]['record_evidence'][0]['quote']='Fabricated content.'
        with tempfile.TemporaryDirectory() as directory:
            path=Path(directory)/'review.docx'
            with self.assertRaisesRegex(ValueError,'Unverified excerpt'):
                render(Review.model_validate(self.data),self.evidence,path)
            self.assertFalse(path.exists())
if __name__=='__main__':unittest.main()