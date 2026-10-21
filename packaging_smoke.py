import io
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
def main():
    from docx import Document
    from config import ROOT
    from engine import demo_report
    from ingest import extract
    from native_host import NativeService, encode_message, read_message
    from openpyxl import Workbook
    import extract_msg
    with tempfile.TemporaryDirectory(prefix='gch-package-') as temp:
        directory = Path(temp)
        service = NativeService(directory)
        try:
            status = service.status()
            if not status['procedures']:
                raise RuntimeError('Bundled procedures were not packaged.')
            for warning in status['procedure_setup_errors']:
                print('PROCEDURE WARNING:', warning)
            sources = service.app.repo.active_procedures()
            if not service.app.repo.search('complaint', [s.id for s, _ in sources]):
                raise RuntimeError('SQLite full-text search failed.')
            demo_report(directory, 'packaging-smoke')
            report = directory/'reports/packaging-smoke/review.docx'
            document = Document(report)
            if not document.tables or len(document.paragraphs) < 10:
                raise RuntimeError('Synthetic Word report is incomplete.')
            example = ROOT/'examples/GCH_Synthetic_Example.docx'
            if not extract(example.name, example.read_bytes()):
                raise RuntimeError('DOCX extraction failed.')
            pdf = ROOT/'procedures/027-P043_K_EN.pdf'
            extract(pdf.name, pdf.read_bytes())
            text = 'Synthetic complaint attachment for portable package validation only.'
            workbook = Workbook()
            workbook.active.append([text])
            spreadsheet = io.BytesIO()
            workbook.save(spreadsheet)
            workbook.close()
            for name, content in (
                ('synthetic.xlsx', spreadsheet.getvalue()),
                ('synthetic.rtf', ('{\\rtf1\\ansi '+text+'}').encode()),
                ('synthetic.eml', ('From: training@example.test\nSubject: Synthetic package check\n\n'+text).encode()),
            ):
                chunks, _ = extract(name, content)
                if not any(text in chunk.text for chunk in chunks):
                    raise RuntimeError(f'Attachment extraction failed: {name}')
            if not callable(extract_msg.openMsg):
                raise RuntimeError('Outlook parser was not packaged.')
        finally:
            service.close()
        origin = 'chrome-extension://' + json.loads((ROOT/'native_config.json').read_text())['extension_id'] + '/'
        command = [sys.executable] if getattr(sys, 'frozen', False) else [sys.executable, str(ROOT/'windows_helper.py')]
        child_env = os.environ.copy()
        child_env.pop('MDT_GPT_API_TOKEN', None)
        child_env['GCH_DATA_DIR'] = str(directory/'native')
        child_env['GCH_DOWNLOADS_DIR'] = str(directory/'downloads')
        result = subprocess.run(command+[origin], input=encode_message({'id':'smoke','method':'status'}),
                                capture_output=True, timeout=60, env=child_env, check=True)
        stream = io.BytesIO(result.stdout)
        response = read_message(stream)
        if not response or response.get('id') != 'smoke' or not response.get('ok') or stream.read():
            raise RuntimeError('Native messaging response was missing or contaminated.')
    print('PASS: bundled assets, SQLite search, DOCX/PDF/XLSX/RTF/EML extraction, Outlook parser import, synthetic Word report, native messaging.')
    return 0