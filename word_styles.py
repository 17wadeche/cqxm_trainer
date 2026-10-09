from math import ceil
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
def document(title: str, subtitle: str):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin, sec.bottom_margin = Inches(.65), Inches(.65)
    sec.left_margin, sec.right_margin = Inches(.65), Inches(.65)
    sec.header_distance, sec.footer_distance = Inches(.25), Inches(.25)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.05
    for name, size in [("Title", 23), ("Subtitle", 11), ("Heading 1", 16), ("Heading 2", 12)]:
        style = doc.styles[name]
        style.font.name, style.font.size = "Calibri", Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.underline = False
        style.paragraph_format.space_before = Pt(10 if name.startswith("Heading") else 0)
        style.paragraph_format.space_after = Pt(6)
        ppr = style.element.find(qn("w:pPr"))
        if ppr is not None:
            for tag in ["w:pBdr", "w:shd"]:
                el = ppr.find(qn(tag))
                if el is not None:
                    ppr.remove(el)
    header = sec.header.paragraphs[0]
    header.text = "GCH CHECK MY WORK  |  TRAINING SUPPORT"
    header.runs[0].font.size = Pt(8)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("Draft for trainer review  |  ").font.size = Pt(8)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.core_properties.title = title
    doc.add_paragraph(title, "Title")
    doc.add_paragraph(subtitle, "Subtitle")
    return doc
def table(doc, headers, rows, widths, font_size=9.5, center_cols=(), allow_split=False):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for col, width in zip(t.columns, widths):
        col.width = Inches(width)
    tblpr = t._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement("w:" + edge)
        for k, v in [("val", "single"), ("sz", "4"), ("color", "D9D9D9")]:
            e.set(qn("w:" + k), v)
        borders.append(e)
    tblpr.append(borders)
    for i, values in enumerate([headers] + list(rows)):
        row = t.rows[0] if i == 0 else t.add_row()
        estimated_lines = max(sum(max(1, ceil(len(line) / max(1, (width-.12)*72/(font_size*.5))))
                                  for line in str(value).split('\n'))
                              for value, width in zip(values, widths))
        if i == 0 or not allow_split or estimated_lines <= 30:
            row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        if i == 0:
            repeat = OxmlElement("w:tblHeader")
            row._tr.get_or_add_trPr().append(repeat)
        for j, (cell, value, width) in enumerate(zip(row.cells, values, widths)):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = str(value)
            props = cell._tc.get_or_add_tcPr()
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), "E7EEF5" if i == 0 else "FFFFFF")
            props.append(shade)
            margins = OxmlElement("w:tcMar")
            for edge, val in [("top", "85"), ("bottom", "85"), ("left", "75"), ("right", "75")]:
                node = OxmlElement("w:" + edge)
                node.set(qn("w:w"), val)
                node.set(qn("w:type"), "dxa")
                margins.append(node)
            props.append(margins)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.0
                if i == 0:
                    p.paragraph_format.keep_with_next = True
                if j in center_cols:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.name, run.font.size = "Calibri", Pt(font_size)
                    run.bold = i == 0
                    run.font.color.rgb = RGBColor(0, 0, 0)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.line_spacing = Pt(1)
    spacer.paragraph_format.space_after = Pt(3)
    spacer.paragraph_format.keep_with_next = True
    return t
def body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        p.add_run(bold_lead).bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    return p