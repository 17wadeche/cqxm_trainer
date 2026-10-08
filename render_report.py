"""Render a reviewed JSON payload into an Anna-style .docx. No network calls."""
from __future__ import annotations

import argparse
from collections import Counter
from pathlib import Path
import re
from docx.shared import Pt

from models import Review, EvidenceManifest, validate_traceability
from word_styles import document, table, body

LABELS = {
    "done_properly": "Done properly", "needs_trainer_review": "Trainer review",
    "potential_gap": "Potential gap", "not_assessable": "Not assessable",
    "not_applicable": "Not applicable",
}

# These limits only decide where prose is displayed. They never reject or
# shorten a finding, and the full text remains in both Word and review.json.
TABLE_LIMITS = {'record_element':120, 'what_was_done':500, 'feedback':600, 'study_action':400}
FIELD_LABELS = {'record_element':'Record element', 'what_was_done':'What was done',
                'feedback':'Feedback for new employee', 'study_action':'Study action'}


def fits_cell(text, limit):
    return len(text)<=limit and text.count('\n')<=6


SHORT_LABELS = {'intake':'Intake', 'narrative':'Narrative', 'identifiers':'Product identifiers',
    'coding':'Event coding', 'harm':'Patient impact', 'gfe':'Good faith effort',
    'investigation':'Investigation', 'cause':'Cause and conclusion', 'reportability':'US reportability',
    'timing':'MDR timing', 'consultation':'Medical consultation', 'consistency':'Event and MDR', 'closure':'Re-closure'}


def compact_references(finding, sources):
    """Keep all cited procedure pages while removing repeated chunks/excerpts."""
    grouped={}
    for ref in finding.procedure_references:
        source=sources[ref.source_id]
        chunk=next(c for c in source.chunks if c.id==ref.chunk_id)
        key=(source.document_id,source.revision)
        page=re.match(r'PDF page (\d+)',chunk.locator)
        grouped.setdefault(key,set()).add(int(page.group(1)) if page else re.sub(r', part \d+$','',chunk.locator))
    labels=[]
    for (docid,revision),locations in grouped.items():
        pages=sorted(v for v in locations if isinstance(v,int))
        other=sorted(v for v in locations if isinstance(v,str))
        location=('p. ' if len(pages)==1 else 'pp. ')+', '.join(map(str,pages)) if pages else ''
        if other:location+='; '+ '; '.join(other) if location else '; '.join(other)
        labels.append(docid+(f' Rev {revision}' if revision else '')+'\n'+location)
    return '\n'.join(labels) or 'Rule not established'


def write_concise_report(*, record_id, stage, findings, procedures, scope, output, note='', run_note=''):
    """Layout-only helper. Live callers validate the complete source review first."""
    doc=document('GCH Training Review','Training feedback | Discuss findings with your trainer')
    doc.styles['Title'].font.size=Pt(20)
    body(doc,f'Record {record_id} | Stage {stage}')
    if note:body(doc,note)
    if procedures:body(doc,'Procedures: '+procedures)
    body(doc,scope)
    counts=Counter(f['assessment'] for f in findings)
    body(doc,' | '.join(f'{counts[key]} {label}' for key,label in [
        ('done_properly','done properly'),('needs_trainer_review','trainer review'),
        ('potential_gap','potential gaps'),('not_assessable','not assessable'),
        ('not_applicable','not applicable')] if counts[key]))
    doc.add_heading('Training review',1)
    rows=[[i,f['area'],f['observation'],f['references'],LABELS[f['assessment']],f['feedback']]
          for i,f in enumerate(findings,1)]
    table(doc,['#','Record element','What was done','Procedure reference','Assessment','Feedback for trainee'],
          rows,[.25,.85,1.35,1.05,.85,2.80],font_size=10,center_cols=(0,),allow_split=True)
    priorities=sorted([(i,f) for i,f in enumerate(findings,1) if f['priority']!='none' and f['action']],
                      key=lambda pair:{'high':0,'medium':1,'low':2}[pair[1]['priority']])[:5]
    if priorities:
        doc.add_heading('First study priorities',1)
        body(doc,'Start with these actions. Other points for discussion are in the review table.')
        table(doc,['Priority','Review area','Action'],
              [[f['priority'].title(),f"Item {i} · {f['area']}",f['action']] for i,f in priorities],
              [.65,1.50,5.05],font_size=10,allow_split=True)
    body(doc,'Trainer ____________________  Date __________  Agreed next step ____________________')
    if run_note:
        p=body(doc,run_note)
        for run in p.runs:run.font.size=Pt(8)
    output.parent.mkdir(parents=True,exist_ok=True)
    doc.save(output)


def render(review: Review, evidence: EvidenceManifest, output: Path, *, detailed=False) -> None:
    if detailed:
        return render_detailed(review,evidence,output)
    validate_traceability(review,evidence)
    sources={s.id:s for s in evidence.sources}
    procedures='; '.join(dict.fromkeys(s.document_id+' Rev '+str(s.revision) for s in evidence.sources if s.kind=='procedure'))
    record_types=', '.join(dict.fromkeys(s.document_id for s in evidence.sources if s.kind=='record')) or 'No record source'
    scope=f'Sources: {record_types}. Assessments apply to the supplied evidence. Missing export content does not prove omitted work; visuals and unavailable sources need trainer confirmation.'
    if not any(s.document_id=='MDR' and s.kind=='record' for s in evidence.sources):
        scope+=' A matching MDR was not supplied.'
    findings=[{'area':SHORT_LABELS.get(f.check_id,f.record_element),'observation':f.what_was_done,
               'references':compact_references(f,sources),'assessment':f.assessment,'feedback':f.feedback,
               'priority':f.priority,'action':f.study_action} for f in review.findings]
    write_concise_report(record_id=evidence.record_id,stage=evidence.record_stage.replace('Displayed saved GCH status: ',''),
        findings=findings,procedures=procedures,scope=scope,output=output,
        note='SYNTHETIC DEMO | No GCH or MDT-GPT connection' if evidence.mode=='demo' else '',
        run_note=f'Review {evidence.review_id} | {evidence.generated_at[:10]} | Full findings and source evidence remain in the saved review files.')


def render_detailed(review: Review, evidence: EvidenceManifest, output: Path) -> None:
    validate_traceability(review, evidence)
    sources = {s.id: s for s in evidence.sources}
    demo = evidence.mode == "demo"
    subtitle = "SYNTHETIC DEMO  |  No GCH or MDT-GPT connection" if demo else "AI assisted training review  |  Trainer confirmation required"
    doc = document("GCH Complaint Training Review", subtitle)
    body(doc, f"Record {evidence.record_id}   |   Stage {evidence.record_stage}")
    body(doc, f"Review {evidence.review_id}   |   Generated {evidence.generated_at}")
    body(doc, review.overall_summary)
    doc.add_heading("Records and procedures reviewed", 1)
    rows = []
    for source in evidence.sources:
        rows.append([source.document_id, source.name, source.revision or "Not applicable"])
    table(doc, ["Source", "Document", "Revision"], rows, [.95, 5.10, 1.15])
    body(doc, "Procedure selection: " + evidence.policy_selection)
    doc.add_heading("Scope and limitations", 1)
    for item in dict.fromkeys(evidence.limitations + review.limitations):
        doc.add_paragraph(item, "List Bullet")
    body(doc, "A missing export field is not proof that work was omitted. Assessments apply only to the supplied evidence and the selected procedure revisions.")

    def citation_label(ref):
        s = sources[ref.source_id]
        c = next(c for c in s.chunks if c.id == ref.chunk_id)
        rev = f", revision {s.revision}" if s.revision else ""
        return f"{s.document_id}{rev}, {c.locator}"

    doc.add_heading("Training review table", 1).paragraph_format.page_break_before = True
    body(doc, "Evidence excerpts follow the table. Trainer review and potential gap findings require a human decision.")
    rows = []
    details = {}
    def cell_text(index, finding, field):
        text=getattr(finding,field)
        if fits_cell(text,TABLE_LIMITS[field]):
            return text
        details.setdefault(index,{})[field]=text
        return f'See item {index} in Detailed review feedback.'

    def item_heading(index, finding):
        return f'Item {index} '+finding.record_element if fits_cell(finding.record_element,120) else f'Item {index}'

    for i, f in enumerate(review.findings, 1):
        refs = "\n".join(citation_label(r) for r in f.procedure_references) or "No applicable rule supplied"
        if not fits_cell(refs,350):
            refs=f'See evidence for item {i}.'
        rows.append([i, cell_text(i,f,'record_element'), cell_text(i,f,'what_was_done'), refs,
                     LABELS[f.assessment], cell_text(i,f,'feedback')])
    table(doc, ["#", "Record element", "What was done", "Procedure reference", "Assessment", "Feedback for new employee"],
          rows, [.25, 1.05, 1.35, 1.20, 1.00, 2.35], font_size=9.5, center_cols=(0,))
    doc.add_heading("Study priorities", 1)
    priorities = sorted([(i,f) for i,f in enumerate(review.findings,1) if f.priority != "none"],
                        key=lambda pair: {"high": 0, "medium": 1, "low": 2}[pair[1].priority])
    if priorities:
        rows = [[f.priority.title(), cell_text(i,f,'record_element'), cell_text(i,f,'study_action')] for i,f in priorities]
        table(doc, ["Priority", "Review area", "Action for trainee and trainer"], rows, [.75, 2.25, 4.20])
    else:
        body(doc, "No study priorities were supplied. This does not establish overall compliance.")

    if details:
        doc.add_heading('Detailed review feedback',1).paragraph_format.page_break_before=True
        body(doc,'Longer explanations are shown in full here. Item numbers match the training review table.')
        for i,fields in sorted(details.items()):
            finding=review.findings[i-1]
            doc.add_heading(item_heading(i,finding),2)
            for field,text in fields.items():
                body(doc,FIELD_LABELS[field],bold_lead=FIELD_LABELS[field])
                body(doc,text).paragraph_format.keep_together=False

    doc.add_heading("Evidence and trainer notes", 1).paragraph_format.page_break_before = True
    for i, f in enumerate(review.findings, 1):
        doc.add_heading(item_heading(i,f), 2)
        refs = f.record_evidence + f.procedure_references
        if not refs:
            body(doc, "No supporting excerpts supplied. See the assessment and limitation above.")
        for ref in refs:
            body(doc, citation_label(ref))
            body(doc, "Excerpt: “" + ref.quote + "”")
    doc.add_heading("Trainer review", 1)
    body(doc, "Trainer name ____________________    Review date ____________________")
    body(doc, "Agreed findings and corrections ______________________________________")
    body(doc, "Follow up learning actions __________________________________________")
    doc.add_heading("Run traceability", 2)
    body(doc, f"Prompt {evidence.prompt_version} | Model {evidence.model_version} | Template {evidence.template_version}")
    body(doc, "Source file and extracted text digests remain in the associated evidence manifest. Excerpt matching verifies text presence only; it does not verify the reasoning or establish that the export is complete.")
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--review", type=Path, required=True)
    parser.add_argument("--evidence", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument('--detailed',action='store_true',help='Include full limitations, narrative detail and source excerpts.')
    args = parser.parse_args()
    review = Review.model_validate_json(args.review.read_text(encoding="utf-8"))
    evidence = EvidenceManifest.model_validate_json(args.evidence.read_text(encoding="utf-8"))
    render(review, evidence, args.output, detailed=args.detailed)
    print(args.output.resolve())


if __name__ == "__main__":
    main()
